"""P3 integration: keyword dataset + internal links (real PostgreSQL).

Runs against the local PostgreSQL (docker container). Skipped when the
database is not reachable, so the suite stays green on CI.

Covers the P3 acceptance:
  - table shapes match spec 46.2-46.4
  - multi-sheet Excel import: Sheet -> cluster, Keyword -> row,
    Volume/KD/CPC parsed
  - repeat import -> upsert, no duplicate records
  - keyword not in dataset -> job still creatable with
    keyword_metrics_available = False (SERP-only mode, section 19.2)
  - keyword in dataset -> keyword_metrics_available = True
  - internal link markers: validate + resolve to markdown links
"""

import uuid
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook
from sqlalchemy import delete, select, text

from app.core.enums import JobStatus
from app.db.models import GenerationJob, InternalLinkRule, Keyword, KeywordCluster
from app.db.session import SessionLocal, check_database, engine
from app.pipeline.steps.keyword_prepare import prepare_keyword
from app.schemas.internal_link import InternalLinkRule as Rule
from app.services import internal_link_service as ils
from app.services.keyword_import import import_workbook

pytestmark = pytest.mark.skipif(
    not check_database(), reason="PostgreSQL not reachable"
)


@pytest.fixture()
def xlsx(tmp_path: Path) -> Path:
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "p3test-attachment"
    ws1.append(["Keyword", "Search Volume", "Keyword Difficulty", "CPC", "Intent"])
    ws1.append(["p3test anxious attachment no contact", 1200, 28.5, 1.25, "Informational"])
    ws1.append(["p3test what is anxious attachment", 2400, 35, 2.10, "Informational"])
    ws1.append(["p3test how to get over anxious attachment", 590, 12, 0.9, "Transactional"])
    ws2 = wb.create_sheet("p3test-avoidance")
    ws2.append(["keyword", "volume", "kd %", "cpc", "intent", "UnknownCol"])
    ws2.append(["p3test avoidant attachment signs", 700, 18.4, 0.95, "Informational", 7])
    path = tmp_path / "semrush.xlsx"
    wb.save(str(path))
    return path


def _wipe_test_data(conn) -> None:
    """Delete only the rows THIS suite creates (p3test-prefixed) — user data
    in the shared tables must survive."""
    conn.execute(
        text(
            "DELETE FROM keywords WHERE cluster_id IN "
            "(SELECT id FROM keyword_clusters WHERE sheet_name LIKE :s)"
        ),
        {"s": "p3test%"},
    )
    conn.execute(
        text("DELETE FROM keyword_clusters WHERE sheet_name LIKE :s"),
        {"s": "p3test%"},
    )
    conn.execute(
        text("DELETE FROM internal_link_rules WHERE marker LIKE :m"),
        {"m": "P3TEST%"},
    )
    conn.execute(
        text("DELETE FROM generation_jobs WHERE keyword LIKE :p"),
        {"p": "p3test%"},
    )


@pytest.fixture(autouse=True)
def _cleanup():
    with engine.begin() as conn:
        _wipe_test_data(conn)
    yield
    with engine.begin() as conn:
        _wipe_test_data(conn)


def _columns(table: str) -> set[str]:
    with engine.connect() as conn:
        rows = conn.execute(
            text(
                "SELECT column_name FROM information_schema.columns "
                "WHERE table_name = :t"
            ),
            {"t": table},
        ).scalars()
        return set(rows)


# ============================================================
# table shapes (spec 46.2-46.4)
# ============================================================
def test_keyword_clusters_table_matches_spec_46_2():
    assert _columns("keyword_clusters") == {
        "id",
        "name",
        "sheet_name",
        "strapi_category_document_id",
        "created_at",
    }


def test_keywords_table_matches_spec_46_3():
    assert _columns("keywords") == {
        "id",
        "cluster_id",
        "keyword",
        "volume",
        "kd",
        "cpc",
        "intent",
        "source",
        "created_at",
    }


def test_internal_link_rules_table_matches_spec_46_4():
    assert _columns("internal_link_rules") == {
        "id",
        "marker",
        "anchor_text",
        "target_url",
        "topic",
        "keywords",
        "active",
        "created_at",
        "updated_at",
    }


def test_keywords_unique_cluster_keyword():
    with SessionLocal() as session:
        cluster = KeywordCluster(name="p3test-uq", sheet_name="p3test-uq-sheet")
        session.add(cluster)
        session.flush()
        session.add(
            Keyword(
                cluster_id=cluster.id,
                keyword="p3test dup kw",
                volume=1,
                source="excel",
            )
        )
        session.flush()
        session.add(
            Keyword(
                cluster_id=cluster.id,
                keyword="p3test dup kw",
                volume=2,
                source="excel",
            )
        )
        with pytest.raises(Exception):
            session.flush()  # IntegrityError on uq_keywords_cluster_keyword
        session.rollback()


# ============================================================
# import: multi-sheet + repeat upsert (P3 acceptance)
# ============================================================
def test_multi_sheet_import_and_repeat_upsert(xlsx: Path):
    with SessionLocal() as session:
        first = import_workbook(session, xlsx, file_name="semrush.xlsx")
        assert first.clusters_created == 2
        assert first.keywords_created == 4
        assert first.keywords_updated == 0

        clusters = {
            c.sheet_name: c
            for c in session.scalars(
                select(KeywordCluster).where(KeywordCluster.sheet_name.like("p3test%"))
            ).all()
        }
        assert set(clusters) == {"p3test-attachment", "p3test-avoidance"}

        kw = session.scalars(
            select(Keyword).where(
                Keyword.keyword == "p3test anxious attachment no contact"
            )
        ).first()
        assert kw is not None
        assert kw.volume == 1200
        assert kw.kd is not None and float(kw.kd) == 28.5
        assert kw.cpc is not None and float(kw.cpc) == 1.25
        assert kw.intent == "Informational"
        assert kw.cluster.sheet_name == "p3test-attachment"
        assert kw.source == "excel"

        # unknown column on sheet 2 warned, never silently guessed
        avoid = clusters["p3test-avoidance"]
        avoid_kw = session.scalars(
            select(Keyword).where(Keyword.cluster_id == avoid.id)
        ).first()
        assert avoid_kw.volume == 700
        sheet_warnings = first.warnings
        assert any("UnknownCol" in w for w in sheet_warnings)

        # ---- repeat import: upsert, no duplicates ----
        second = import_workbook(session, xlsx, file_name="semrush.xlsx")
        assert second.clusters_created == 0
        assert second.keywords_created == 0
        assert second.keywords_updated == 4

        count = session.scalar(
            text(
                "SELECT COUNT(*) FROM keywords WHERE cluster_id IN "
                "(SELECT id FROM keyword_clusters WHERE sheet_name LIKE :s)"
            ),
            {"s": "p3test%"},
        )
        assert count == 4

        cluster_count = session.scalar(
            text("SELECT COUNT(*) FROM keyword_clusters WHERE sheet_name LIKE :s"),
            {"s": "p3test%"},
        )
        assert cluster_count == 2


# ============================================================
# SERP-only mode (section 19.2, P3 acceptance)
# ============================================================
def _make_job(keyword: str) -> GenerationJob:
    with SessionLocal() as session:
        job = GenerationJob(
            id=uuid.uuid4(),
            keyword=keyword,
            language="en",
            market="US",
            strategy="auto",
            status=JobStatus.QUEUED.value,
            current_step="queued",
        )
        session.add(job)
        session.commit()
        session.refresh(job)
        return job


def test_keyword_in_dataset_sets_metrics_available(tmp_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "p3test-p3"
    ws.append(["Keyword", "Volume", "KD", "CPC", "Intent"])
    ws.append(["p3test in-dataset keyword", 4200, 21, 1.5, "Informational"])
    path = tmp_path / "p3.xlsx"
    wb.save(str(path))
    with SessionLocal() as session:
        import_workbook(session, path, file_name="p3.xlsx")

    job = _make_job("p3test in-dataset keyword")
    with SessionLocal() as session:
        dbjob = session.get(GenerationJob, job.id)
        metrics = prepare_keyword(session, dbjob)
        assert metrics is not None
        assert metrics.volume == 4200
        assert dbjob.keyword_metrics_available is True
        assert dbjob.status == JobStatus.SERP_SEARCHING.value


def test_keyword_not_in_dataset_is_serp_only(xlsx: Path):
    with SessionLocal() as session:
        import_workbook(session, xlsx, file_name="semrush.xlsx")

    # job is still creatable for a dataset-unknown keyword (section 19.2)
    job = _make_job("p3test-unknown keyword not in dataset")
    with SessionLocal() as session:
        dbjob = session.get(GenerationJob, job.id)
        assert dbjob.status == JobStatus.QUEUED.value
        metrics = prepare_keyword(session, dbjob)
        assert metrics is None
        assert dbjob.keyword_metrics_available is False
        # job continues, not failed
        assert dbjob.status == JobStatus.SERP_SEARCHING.value


# ============================================================
# internal links (section 21)
# ============================================================
def test_internal_link_roundtrip():
    with SessionLocal() as session:
        ils.upsert_rule(
            session,
            Rule(
                marker="P3TEST_ATTACHMENT_TEST",
                anchor_text="what is anxious attachment",
                target_url="https://example.com/what-is-anxious-attachment",
                topic="attachment",
                keywords=["anxious attachment"],
            ),
        )
        session.commit()

        rules = session.scalars(
            select(InternalLinkRule).where(
                InternalLinkRule.marker.like("P3TEST%")
            )
        ).all()
        assert len(rules) == 1
        assert rules[0].keywords == ["anxious attachment"]

        md = (
            "## When to reach out\n\n"
            "Read [[INTERNAL_LINK:P3TEST_ATTACHMENT_TEST]] first, then "
            "[[INTERNAL_LINK:P3TEST_GHOST]] for more."
        )
        rendered, resolved, validation = ils.resolve_markers(session, md)
        assert (
            "[what is anxious attachment]"
            "(https://example.com/what-is-anxious-attachment)" in rendered
        )
        assert validation.valid is False
        assert validation.unknown_markers == ["P3TEST_GHOST"]
        assert len(resolved) == 1

        # duplicate marker rejected at DB level (unique)
        with pytest.raises(Exception):
            session.add(
                InternalLinkRule(
                    marker="P3TEST_ATTACHMENT_TEST",
                    target_url="https://example.com/x",
                )
            )
            session.commit()
        session.rollback()
