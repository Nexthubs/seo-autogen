"""P3 unit: SEMrush Excel import (spec sections 19, 20).

Workbooks are built in-memory with openpyxl — no network. DB upsert is
exercised against SQLite so the unit suite stays hermetic; the real-
PostgreSQL roundtrip lives in tests/integration/test_keyword_dataset.py.
"""

from pathlib import Path

import pytest
from openpyxl import Workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session

from app.db.base import Base
from app.db.models import Keyword
from app.services.keyword_import import (
    _canonical_header,
    parse_workbook_with_rows,
)


@pytest.fixture()
def db():
    """In-memory SQLite so the unit suite stays hermetic (no network/PG)."""
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    with Session(engine) as session:
        yield session
    engine.dispose()


def _write_workbook(path: Path) -> Path:
    wb = Workbook()
    # Sheet 1: alias headers + an unknown column.
    ws1 = wb.active
    ws1.title = "Attachment"
    ws1.append(["Keyword", "Search Volume", "Keyword Difficulty", "CPC", "Intent", "Competitors"])
    ws1.append(["anxious attachment no contact", 1200, 28.5, 1.25, "Informational", 42])
    ws1.append(["what is anxious attachment", 2400, 35, 2.10, "Informational", 38])
    ws1.append(["", 999, 9.9, 0.5, "Transactional", 1])  # empty keyword
    ws1.append(["bad volume keyword", "not-a-number", 12, 0.8, "Commercial", 5])
    # Sheet 2: lowercase / other aliases (section 20).
    ws2 = wb.create_sheet("Avoidance")
    ws2.append(["keyword", "volume", "kd %", "cpc", "search intent"])
    ws2.append(["avoidant attachment signs", 700, "18.4", "$0.95", "Informational"])
    wb.save(str(path))
    return path


# ============================================================
# header alias mapping (section 20)
# ============================================================
@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("Keyword", "keyword"),
        ("keyword", "keyword"),
        ("Volume", "volume"),
        ("Search Volume", "volume"),
        ("KD", "kd"),
        ("Keyword Difficulty", "kd"),
        ("KD %", "kd"),
        ("CPC", "cpc"),
        ("Intent", "intent"),
        ("Search Intent", "intent"),
        ("Competitors", None),
        ("Random", None),
        (None, None),
    ],
)
def test_header_aliases(raw, expected):
    assert _canonical_header(raw) == expected


# ============================================================
# workbook parsing
# ============================================================
def test_parse_multi_sheet_workbook(tmp_path):
    path = _write_workbook(tmp_path / "semrush.xlsx")
    shells, rows_by_sheet = parse_workbook_with_rows(path)

    assert [s.sheet_name for s in shells] == ["Attachment", "Avoidance"]
    # Sheet 1: unknown column warned, never guessed.
    attachment = rows_by_sheet["Attachment"]
    assert len(attachment) == 3  # empty-keyword row skipped
    assert attachment[0].keyword == "anxious attachment no contact"
    assert attachment[0].volume == 1200
    assert attachment[0].kd is not None and float(attachment[0].kd) == 28.5
    assert attachment[0].cpc is not None and float(attachment[0].cpc) == 1.25
    assert attachment[0].intent == "Informational"
    # Unparseable volume -> kept with None + warning.
    bad = attachment[2]
    assert bad.volume is None
    shell_warnings = " | ".join(
        w for s in shells for w in s.warnings if s.sheet_name == "Attachment"
    )
    assert "Competitors" in shell_warnings
    assert "not-a-number" in shell_warnings
    assert "empty keyword" in shell_warnings

    # Sheet 2: alias headers resolved.
    avoidance = rows_by_sheet["Avoidance"]
    assert avoidance[0].keyword == "avoidant attachment signs"
    assert avoidance[0].volume == 700
    assert float(avoidance[0].kd) == 18.4
    assert float(avoidance[0].cpc) == 0.95
    assert avoidance[0].intent == "Informational"


def test_parse_missing_keyword_column_warns(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "NoKeyword"
    ws.append(["Volume", "KD"])
    ws.append([10, 5])
    path = tmp_path / "nokeyword.xlsx"
    wb.save(str(path))
    shells, rows_by_sheet = parse_workbook_with_rows(path)
    assert rows_by_sheet["NoKeyword"] == []
    assert any("Keyword" in w for w in shells[0].warnings)


# ============================================================
# upsert behaviour (section 20: repeat import -> no duplicates)
# ============================================================
def test_import_then_reimport_is_upsert(tmp_path, db: Session):
    from app.services.keyword_import import import_workbook

    path = _write_workbook(tmp_path / "semrush.xlsx")
    first = import_workbook(db, path, file_name="semrush.xlsx")
    assert first.clusters_created == 2
    assert first.keywords_created == 4  # 3 Attachment + 1 Avoidance
    assert first.keywords_updated == 0

    # Bump one metric, re-import.
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Attachment"
    ws1.append(["Keyword", "Search Volume", "Keyword Difficulty", "CPC", "Intent"])
    ws1.append(["anxious attachment no contact", 9999, 10, 0.5, "Informational"])
    ws1.append(["what is anxious attachment", 2400, 35, 2.10, "Informational"])
    ws1.append(["new keyword row", 10, 5, 0.1, "Navigational"])
    ws2 = wb.create_sheet("Avoidance")
    ws2.append(["keyword", "volume", "kd %", "cpc", "search intent"])
    ws2.append(["avoidant attachment signs", 701, 18.4, 0.95, "Informational"])
    path2 = tmp_path / "semrush_v2.xlsx"
    wb.save(str(path2))

    second = import_workbook(db, path2, file_name="semrush_v2.xlsx")
    assert second.clusters_created == 0
    assert second.keywords_created == 1
    assert second.keywords_updated == 3

    count = len(db.scalars(select(Keyword)).all())
    assert count == 5  # 4 original + 1 new, nothing duplicated

    updated = db.scalars(
        select(Keyword).where(Keyword.keyword == "anxious attachment no contact")
    ).first()
    assert updated.volume == 9999
    assert updated.source == "excel"


def test_import_sheet_without_keyword_column_yields_no_rows(tmp_path, db: Session):
    from app.services.keyword_import import import_workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "NoKeyword"
    ws.append(["Volume", "KD"])
    ws.append([10, 5])
    path = tmp_path / "nokeyword.xlsx"
    wb.save(str(path))
    report = import_workbook(db, path)
    assert report.clusters_created == 1
    assert report.keywords_created == 0
    assert report.keywords_updated == 0
    assert any("Keyword" in w for w in report.warnings)
